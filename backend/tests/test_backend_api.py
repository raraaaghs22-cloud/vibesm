"""Backend API tests for VIBESMAI (new schema)."""
import io
import time
import uuid
import pytest
from openpyxl import load_workbook


# ---------- Health ----------
def test_root(anon_client, base_url):
    r = anon_client.get(f"{base_url}/api/")
    assert r.status_code == 200
    assert "VIBESMAI" in r.json().get("message", "")


# ---------- Auth / RBAC ----------
class TestAuth:
    def test_me_unauth(self, anon_client, base_url):
        assert anon_client.get(f"{base_url}/api/auth/me").status_code == 401

    def test_me_admin(self, admin_client, base_url):
        r = admin_client.get(f"{base_url}/api/auth/me")
        assert r.status_code == 200
        d = r.json()
        assert d["is_admin"] is True
        assert d["email"].lower() == "raraaaghs22@gmail.com"

    def test_me_nonadmin(self, nonadmin_client, base_url):
        r = nonadmin_client.get(f"{base_url}/api/auth/me")
        assert r.status_code == 200
        assert r.json()["is_admin"] is False

    @pytest.mark.parametrize("ep", [
        "/api/admin/submissions", "/api/admin/stats", "/api/admin/settings"
    ])
    def test_admin_endpoints_401_unauth(self, anon_client, base_url, ep):
        assert anon_client.get(f"{base_url}{ep}").status_code == 401

    @pytest.mark.parametrize("ep", [
        "/api/admin/submissions", "/api/admin/stats", "/api/admin/settings"
    ])
    def test_admin_endpoints_403_nonadmin(self, nonadmin_client, base_url, ep):
        assert nonadmin_client.get(f"{base_url}{ep}").status_code == 403


# ---------- Submissions ----------
class TestSubmissions:
    def test_submit_youtube_video_link_field(self, anon_client, base_url, db):
        r = anon_client.post(f"{base_url}/api/submissions", json={
            "full_name": "TEST_YT One", "class_name": "XI 3",
            "attendance_number": 7, "video_link": "https://www.youtube.com/watch?v=dQw4w9WgXcQ"})
        assert r.status_code == 200
        body = r.json()
        assert body["ok"] is True
        assert body["message"] == "Your video assignment link has been successfully submitted / Tugas link video Anda berhasil dikirimkan."
        rec = db.submissions.find_one({"full_name": "TEST_YT One"})
        assert rec is not None
        assert rec["platform"] == "youtube"
        assert rec["video_link"].startswith("https://")
        assert rec["status"] in ("pending", "processing", "draft", "failed")

    def test_submit_legacy_video_url_accepted(self, anon_client, base_url, db):
        r = anon_client.post(f"{base_url}/api/submissions", json={
            "full_name": "TEST_Legacy", "class_name": "XI 1",
            "attendance_number": 2, "video_url": "https://www.tiktok.com/@u/video/123"})
        assert r.status_code == 200
        rec = db.submissions.find_one({"full_name": "TEST_Legacy"})
        assert rec["platform"] == "tiktok"

    def test_submit_invalid_platform(self, anon_client, base_url):
        r = anon_client.post(f"{base_url}/api/submissions", json={
            "full_name": "TEST_Bad", "class_name": "XI 1",
            "attendance_number": 1, "video_link": "https://example.com/x"})
        assert r.status_code == 422

    def test_submit_invalid_class(self, anon_client, base_url):
        r = anon_client.post(f"{base_url}/api/submissions", json={
            "full_name": "TEST_Bad2", "class_name": "XI 13",
            "attendance_number": 1, "video_link": "https://youtube.com/watch?v=a"})
        assert r.status_code == 422

    def test_submit_short_name(self, anon_client, base_url):
        r = anon_client.post(f"{base_url}/api/submissions", json={
            "full_name": "A", "class_name": "XI 1",
            "attendance_number": 1, "video_link": "https://youtube.com/watch?v=a"})
        assert r.status_code == 422


# ---------- Admin list / filters / stats ----------
def _seed(db, **kw):
    sid = str(uuid.uuid4())
    base = {
        "id": sid, "full_name": "TEST_X", "class_name": "XI 7", "attendance_number": 10,
        "video_link": "https://youtu.be/abc", "platform": "youtube", "status": "draft",
        "ai_score": 80, "ai_letter_grade": "B", "ai_strengths": "s", "ai_weaknesses": "w",
        "content_score": 80, "delivery_score": 80, "technical_score": 80,
        "final_score": 80.0, "final_grade": "B", "manually_edited": False,
        "teacher_notes": "", "created_at": "2026-01-01T00:00:00+00:00",
    }
    base.update(kw)
    db.submissions.insert_one(base)
    return sid


class TestAdminList:
    def test_list_ok(self, admin_client, base_url):
        r = admin_client.get(f"{base_url}/api/admin/submissions")
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_filter_class(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_FltC", class_name="XI 2")
        r = admin_client.get(f"{base_url}/api/admin/submissions", params={"class_name": "XI 2"})
        assert r.status_code == 200
        assert all(s["class_name"] == "XI 2" for s in r.json())

    def test_filter_status_pending_includes_processing(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_Pend", status="pending", class_name="XI 8")
        _seed(db, full_name="TEST_Proc", status="processing", class_name="XI 8")
        r = admin_client.get(f"{base_url}/api/admin/submissions", params={"status": "pending", "class_name": "XI 8"})
        assert r.status_code == 200
        names = {s["full_name"] for s in r.json()}
        assert "TEST_Pend" in names and "TEST_Proc" in names

    def test_filter_q_name(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_UniqueFilterName", class_name="XI 9")
        r = admin_client.get(f"{base_url}/api/admin/submissions", params={"q": "UniqueFilterName"})
        assert r.status_code == 200
        assert any(s["full_name"] == "TEST_UniqueFilterName" for s in r.json())

    def test_stats_keys(self, admin_client, base_url):
        r = admin_client.get(f"{base_url}/api/admin/stats")
        assert r.status_code == 200
        d = r.json()
        for k in ("total", "draft", "final", "pending", "failed", "avg_score", "per_class"):
            assert k in d
        assert set(d["per_class"].keys()) >= {f"XI {i}" for i in range(1, 13)}


# ---------- PATCH ----------
class TestPatch:
    def test_patch_final_score_sets_grade(self, admin_client, base_url, db):
        sid = _seed(db, full_name="TEST_PatchFS", class_name="XI 6")
        r = admin_client.patch(f"{base_url}/api/admin/submissions/{sid}", json={"final_score": 88})
        assert r.status_code == 200
        d = r.json()
        assert d["final_score"] == 88.0
        assert d["final_grade"] == "B"
        assert d["manually_edited"] is True

    def test_patch_rubric_weighted(self, admin_client, base_url, db):
        sid = _seed(db, full_name="TEST_PatchRub", class_name="XI 6")
        r = admin_client.patch(f"{base_url}/api/admin/submissions/{sid}",
                               json={"content_score": 90, "delivery_score": 80, "technical_score": 70})
        assert r.status_code == 200
        d = r.json()
        # 90*0.5 + 80*0.3 + 70*0.2 = 83 -> B
        assert d["final_score"] == 83.0
        assert d["final_grade"] == "B"

    def test_patch_status_final_sets_status(self, admin_client, base_url, db):
        sid = _seed(db, full_name="TEST_PatchFinal", class_name="XI 6", final_score=90.0, final_grade="A")
        r = admin_client.patch(f"{base_url}/api/admin/submissions/{sid}", json={"status": "final"})
        assert r.status_code == 200
        assert r.json()["status"] == "final"

    def test_patch_final_without_score_422(self, admin_client, base_url, db):
        sid = _seed(db, full_name="TEST_PatchNoScore", class_name="XI 6",
                    final_score=None, final_grade=None, ai_score=None,
                    status="pending")
        r = admin_client.patch(f"{base_url}/api/admin/submissions/{sid}", json={"status": "final"})
        assert r.status_code == 422


# ---------- Bulk / Delete ----------
class TestBulkDelete:
    def test_bulk_status_final(self, admin_client, base_url, db):
        ids = [_seed(db, full_name=f"TEST_Bulk{i}", class_name="XI 10", final_score=80.0, final_grade="B")
               for i in range(2)]
        r = admin_client.post(f"{base_url}/api/admin/bulk-status", json={"ids": ids, "status": "final"})
        assert r.status_code == 200
        assert r.json()["updated"] == 2
        for sid in ids:
            assert db.submissions.find_one({"id": sid})["status"] == "final"

    def test_delete_then_404(self, admin_client, base_url, db):
        sid = _seed(db, full_name="TEST_Del", class_name="XI 11")
        r = admin_client.delete(f"{base_url}/api/admin/submissions/{sid}")
        assert r.status_code == 200
        r2 = admin_client.delete(f"{base_url}/api/admin/submissions/{sid}")
        assert r2.status_code == 404


# ---------- Export ----------
class TestExport:
    def test_export_csv(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_CSV", class_name="XI 1", status="draft")
        r = admin_client.get(f"{base_url}/api/admin/export", params={"format": "csv"})
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        text = r.text
        assert "Nama Siswa" in text
        assert "Skor AI" in text
        assert "Letter Grade" in text
        # Status labels
        assert "Draft" in text or "Final" in text

    def test_export_xlsx_all_classes(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_Xlsx", class_name="XI 1", status="final")
        r = admin_client.get(f"{base_url}/api/admin/export", params={"format": "xlsx"})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Semua Kelas" in wb.sheetnames
        assert "XI 1" in wb.sheetnames

    def test_export_xlsx_one_class(self, admin_client, base_url, db):
        _seed(db, full_name="TEST_XlsxOne", class_name="XI 3")
        r = admin_client.get(f"{base_url}/api/admin/export",
                             params={"format": "xlsx", "class_name": "XI 3"})
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Kelas XI 3" in wb.sheetnames
        assert "attachment" in r.headers.get("content-disposition", "")


# ---------- Public results ----------
class TestPublicResults:
    def test_public_results_disabled_403(self, admin_client, anon_client, base_url):
        admin_client.put(f"{base_url}/api/admin/settings", json={"results_public": False})
        r = anon_client.get(f"{base_url}/api/public/results",
                            params={"class_name": "XI 1", "attendance_number": 1})
        assert r.status_code == 403

    def test_public_results_enabled_only_final_exposed(self, admin_client, anon_client, base_url, db):
        # Seed one final and one draft
        _seed(db, full_name="TEST_PubFinal", class_name="XI 12", attendance_number=11,
              status="final", final_score=92.0, final_grade="A")
        _seed(db, full_name="TEST_PubDraft", class_name="XI 12", attendance_number=12,
              status="draft", final_score=70.0, final_grade="C")
        try:
            admin_client.put(f"{base_url}/api/admin/settings", json={"results_public": True})
            # Final visible
            r1 = anon_client.get(f"{base_url}/api/public/results",
                                 params={"class_name": "XI 12", "attendance_number": 11})
            assert r1.status_code == 200
            d1 = r1.json()
            assert len(d1) >= 1
            final = next(x for x in d1 if x["full_name"] == "TEST_PubFinal")
            assert final["is_final"] is True
            assert final["final_score"] == 92.0
            assert final["final_grade"] == "A"
            # Draft score hidden
            r2 = anon_client.get(f"{base_url}/api/public/results",
                                 params={"class_name": "XI 12", "attendance_number": 12})
            assert r2.status_code == 200
            d2 = r2.json()
            draft = next(x for x in d2 if x["full_name"] == "TEST_PubDraft")
            assert draft["is_final"] is False
            assert draft["final_score"] is None
            assert draft["final_grade"] is None
            # No sensitive fields leaked
            for forbidden in ("ai_strengths", "ai_weaknesses", "teacher_notes",
                              "content_score", "delivery_score", "technical_score"):
                assert forbidden not in final
                assert forbidden not in draft
        finally:
            admin_client.put(f"{base_url}/api/admin/settings", json={"results_public": False})


# ---------- Seeded AI-graded sample verifications ----------
class TestSeededAISamples:
    """Verify N/A privacy semantics on a seeded failed-extraction submission (self-seeded, not legacy)."""

    def test_seeded_ig_private_na(self, db):
        sid = _seed(db, full_name="TEST_IG_Privat", class_name="XI 5", attendance_number=22,
                    platform="instagram", video_link="https://www.instagram.com/reel/XXXXXXX/",
                    status="draft", ai_score=0, ai_letter_grade="N/A",
                    ai_strengths="-", ai_weaknesses="Sistem tidak dapat menonton video karena tautan diprivasi atau diblokir platform. Silakan klik tautan dan nilai secara manual.",
                    final_score=0.0, final_grade="N/A",
                    content_score=None, delivery_score=None, technical_score=None,
                    extraction_ok=False, ai_input="Video tidak dapat diekstrak karena tautan diprivasi atau diblokir platform.")
        rec = db.submissions.find_one({"id": sid})
        assert rec["ai_score"] == 0
        assert rec["ai_letter_grade"] == "N/A"
        assert rec["final_grade"] == "N/A"
        assert rec["extraction_ok"] is False
        assert "privasi" in rec["ai_weaknesses"].lower()
        assert rec["content_score"] is None
        assert rec["delivery_score"] is None
        assert rec["technical_score"] is None


# ---------- Live AI grading (slow, optional) ----------
@pytest.mark.slow
def test_live_ai_grading_youtube(anon_client, base_url, db):
    r = anon_client.post(f"{base_url}/api/submissions", json={
        "full_name": "TEST_LiveAI", "class_name": "XI 11",
        "attendance_number": 40,
        "video_link": "https://www.youtube.com/watch?v=dQw4w9WgXcQ"})
    assert r.status_code == 200
    end = time.time() + 90
    status = None
    while time.time() < end:
        rec = db.submissions.find_one({"full_name": "TEST_LiveAI"})
        status = rec and rec.get("status")
        if status in ("draft", "failed"):
            break
        time.sleep(5)
    print(f"Live AI final status: {status}")
    assert status in ("draft", "failed", "processing", "pending")
