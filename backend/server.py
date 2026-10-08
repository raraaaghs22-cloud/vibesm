from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, BackgroundTasks, Depends, Query
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import asyncio
import heapq
import itertools
import io
import re
import csv
import json
import html
import uuid
import logging
import httpx
from pathlib import Path
from pydantic import BaseModel, Field, field_validator
from typing import List, Optional
from datetime import datetime, timezone, timedelta
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import tempfile
import shutil
import yt_dlp
import imageio_ffmpeg
from emergentintegrations.llm.chat import LlmChat, UserMessage, FileContentWithMimeType

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

client = AsyncIOMotorClient(os.environ['MONGO_URL'])
db = client[os.environ['DB_NAME']]
EMERGENT_LLM_KEY = os.environ['EMERGENT_LLM_KEY']
ADMIN_EMAIL = os.environ.get('ADMIN_EMAIL', '').strip().lower()

app = FastAPI()
api = APIRouter(prefix="/api")
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

CLASSES = [f"XI {i}" for i in range(1, 13)]
PLATFORM_PATTERNS = {
    "youtube": r"(^|\.)(youtube\.com|youtu\.be)$",
    "tiktok": r"(^|\.)tiktok\.com$",
    "instagram": r"(^|\.)(instagram\.com|instagr\.am)$",
    "facebook": r"(^|\.)(facebook\.com|fb\.watch|fb\.com)$",
}
EXTRACTION_FAILED = "Video tidak dapat diekstrak karena tautan diprivasi atau diblokir platform."
PRIVACY_WEAKNESS = "Sistem tidak dapat menonton video karena tautan diprivasi atau diblokir platform. Silakan klik tautan dan nilai secara manual."
MAX_VIDEO_BYTES = 50 * 1024 * 1024


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def detect_platform(url: str) -> Optional[str]:
    try:
        host = (urlparse(url).hostname or "").lower()
    except ValueError:
        return None
    for name, pat in PLATFORM_PATTERNS.items():
        if re.search(pat, host):
            return name
    return None


def letter_grade(score: float) -> str:
    if score > 90:
        return "A"
    if score >= 80:
        return "B"
    if score >= 70:
        return "C"
    return "D"


def has_readable_text(meta: dict) -> bool:
    texts = [meta.get("oembed", {}).get("title", "")]
    texts += [str(v) for k, v in meta.get("page", {}).items() if k in ("title", "description", "full_description", "keywords")]
    blob = " ".join(t for t in texts if t).strip()
    if len(blob) < 15:
        return False
    return not re.search(r"access denied|forbidden|\b40[134]\b|login|log in|masuk|not available|tidak tersedia|private|sorry|unavailable",
                         blob, re.I)


def weighted(content: float, delivery: float, technical: float) -> float:
    return round(content * 0.5 + delivery * 0.3 + technical * 0.2, 1)


# ---------- Models ----------
class SubmissionIn(BaseModel):
    full_name: str = Field(min_length=2, max_length=120)
    class_name: str
    attendance_number: int = Field(ge=1, le=60)
    video_link: Optional[str] = Field(default=None, max_length=1000)
    video_url: Optional[str] = Field(default=None, max_length=1000)  # legacy alias

    @field_validator("class_name")
    @classmethod
    def check_class(cls, v):
        if v not in CLASSES:
            raise ValueError("Invalid class")
        return v

    @field_validator("full_name", "video_link", "video_url")
    @classmethod
    def strip(cls, v):
        return v.strip() if isinstance(v, str) else v


class SubmissionUpdate(BaseModel):
    content_score: Optional[float] = Field(default=None, ge=0, le=100)
    delivery_score: Optional[float] = Field(default=None, ge=0, le=100)
    technical_score: Optional[float] = Field(default=None, ge=0, le=100)
    final_score: Optional[float] = Field(default=None, ge=0, le=100)
    ai_strengths: Optional[str] = None
    ai_weaknesses: Optional[str] = None
    teacher_notes: Optional[str] = None
    student_feedback: Optional[str] = Field(default=None, max_length=2000)
    status: Optional[str] = Field(default=None, pattern="^(draft|final)$")


class FeedbackIn(BaseModel):
    ai_strengths: Optional[str] = None
    ai_weaknesses: Optional[str] = None
    final_score: Optional[float] = Field(default=None, ge=0, le=100)


class BulkStatus(BaseModel):
    ids: List[str]
    status: str = Field(pattern="^(draft|final)$")


class SettingsIn(BaseModel):
    results_public: bool


# ---------- Auth ----------
async def get_admin_email() -> Optional[str]:
    doc = await db.settings.find_one({"key": "admin"}, {"_id": 0})
    return doc["email"] if doc else None


async def current_user(request: Request) -> dict:
    token = request.cookies.get("session_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    sess = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not sess:
        raise HTTPException(401, "Invalid session")
    exp = sess["expires_at"]
    if isinstance(exp, str):
        exp = datetime.fromisoformat(exp)
    if exp.tzinfo is None:
        exp = exp.replace(tzinfo=timezone.utc)
    if exp < datetime.now(timezone.utc):
        raise HTTPException(401, "Session expired")
    user = await db.users.find_one({"user_id": sess["user_id"]}, {"_id": 0})
    if not user:
        raise HTTPException(401, "User not found")
    admin_email = await get_admin_email()
    user["is_admin"] = bool(admin_email) and user["email"].lower() == admin_email
    return user


async def require_admin(user: dict = Depends(current_user)) -> dict:
    if not user["is_admin"]:
        raise HTTPException(403, "Admin access only")
    return user


@api.post("/auth/session")
async def create_session(request: Request, response: Response):
    session_id = request.headers.get("X-Session-ID")
    if not session_id:
        raise HTTPException(400, "Missing session id")
    async with httpx.AsyncClient(timeout=20) as hc:
        r = await hc.get("https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
                         headers={"X-Session-ID": session_id})
    if r.status_code != 200:
        raise HTTPException(401, "Invalid session id")
    data = r.json()
    email = data["email"].lower()
    existing = await db.users.find_one({"email": email}, {"_id": 0})
    if existing:
        user_id = existing["user_id"]
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": data.get("name"), "picture": data.get("picture")}})
    else:
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        await db.users.insert_one({"user_id": user_id, "email": email, "name": data.get("name"),
                                   "picture": data.get("picture"), "created_at": now_iso()})
    # First account to log in becomes admin (only if no admin assigned yet)
    await db.settings.update_one({"key": "admin"}, {"$setOnInsert": {"key": "admin", "email": email}}, upsert=True)
    await db.user_sessions.insert_one({
        "user_id": user_id, "session_token": data["session_token"],
        "expires_at": datetime.now(timezone.utc) + timedelta(days=7), "created_at": now_iso()})
    response.set_cookie("session_token", data["session_token"], httponly=True, secure=True,
                        samesite="none", path="/", max_age=7 * 24 * 3600)
    admin_email = await get_admin_email()
    return {"user_id": user_id, "email": email, "name": data.get("name"), "picture": data.get("picture"),
            "is_admin": email == admin_email}


@api.get("/auth/me")
async def me(user: dict = Depends(current_user)):
    return user


@api.post("/auth/logout")
async def logout(request: Request, response: Response):
    token = request.cookies.get("session_token")
    if token:
        await db.user_sessions.delete_one({"session_token": token})
    response.delete_cookie("session_token", path="/", secure=True, samesite="none")
    return {"ok": True}


# ---------- Metadata extraction ----------
def meta_content(page: str, *names) -> Optional[str]:
    for n in names:
        m = re.search(rf'<meta[^>]+(?:property|name|itemprop)=["\']{re.escape(n)}["\'][^>]*content=["\']([^"\']*)["\']', page, re.I) \
            or re.search(rf'<meta[^>]+content=["\']([^"\']*)["\'][^>]*(?:property|name|itemprop)=["\']{re.escape(n)}["\']', page, re.I)
        if m and m.group(1).strip():
            return html.unescape(m.group(1).strip())
    return None


async def fetch_metadata(url: str, platform: str) -> dict:
    meta = {"platform": platform, "url": url}
    path = urlparse(url).path.lower()
    if platform == "youtube" and "/shorts/" in path:
        meta["url_hint"] = "YouTube Shorts URL (format vertikal 9:16, maks 60 detik kecuali akun tertentu)"
    if platform == "instagram" and "/reel" in path:
        meta["url_hint"] = "Instagram Reel URL (biasanya vertikal 9:16)"
    if platform == "facebook" and ("/reel" in path or "fb.watch" in url):
        meta["url_hint"] = "Facebook Reel/Watch URL"
    if platform == "tiktok":
        meta["url_hint"] = "TikTok video (biasanya vertikal 9:16)"
    headers = {"User-Agent": "Mozilla/5.0 (compatible; facebookexternalhit/1.1; +http://www.facebook.com/externalhit_uatext.php)",
               "Accept-Language": "id,en;q=0.8"}
    oembed_url = {"youtube": "https://www.youtube.com/oembed", "tiktok": "https://www.tiktok.com/oembed"}.get(platform)
    async with httpx.AsyncClient(timeout=15, follow_redirects=True, headers=headers) as hc:
        if oembed_url:
            try:
                r = await hc.get(oembed_url, params={"url": url, "format": "json"})
                if r.status_code == 200:
                    o = r.json()
                    meta["oembed"] = {k: o.get(k) for k in ("title", "author_name", "author_url", "width", "height",
                                                            "thumbnail_width", "thumbnail_height") if o.get(k) is not None}
            except Exception as e:
                logger.info(f"oEmbed failed: {e}")
        try:
            r = await hc.get(url)
            if r.status_code == 200:
                page = r.text[:600000]
                meta["page"] = {k: v for k, v in {
                    "title": meta_content(page, "og:title", "twitter:title", "title"),
                    "description": meta_content(page, "og:description", "description", "twitter:description"),
                    "video_width": meta_content(page, "og:video:width"),
                    "video_height": meta_content(page, "og:video:height"),
                    "duration_iso": meta_content(page, "duration"),
                    "keywords": meta_content(page, "keywords"),
                }.items() if v}
                if platform == "youtube":
                    m = re.search(r'"lengthSeconds":"(\d+)"', page)
                    if m:
                        meta["page"]["duration_seconds"] = int(m.group(1))
                    m = re.search(r'"shortDescription":"((?:[^"\\]|\\.)*)"', page)
                    if m:
                        try:
                            meta["page"]["full_description"] = json.loads(f'"{m.group(1)}"')[:3000]
                        except Exception:
                            pass
                if platform == "tiktok":
                    m = re.search(r'"duration":(\d+)', page)
                    if m:
                        meta["page"]["duration_seconds"] = int(m.group(1))
            else:
                meta["page_status"] = r.status_code
        except Exception as e:
            meta["page_error"] = str(e)[:200]
    return meta


# ---------- LLM helper ----------
# The LLM gateway limits concurrent requests, so calls are queued per purpose and retried with backoff.
class PriorityLock:
    """Single-slot lock where lower priority numbers are served first (teacher actions before background grading)."""

    def __init__(self):
        self._locked = False
        self._waiters = []
        self._seq = itertools.count()

    async def acquire(self, prio: int):
        if not self._locked:
            self._locked = True
            return
        fut = asyncio.get_running_loop().create_future()
        heapq.heappush(self._waiters, (prio, next(self._seq), fut))
        try:
            await fut
        except asyncio.CancelledError:
            if fut.done() and not fut.cancelled():
                self.release()
            raise

    def release(self):
        while self._waiters:
            _, _, fut = heapq.heappop(self._waiters)
            if not fut.done():
                fut.set_result(True)
                return
        self._locked = False


LLM_LOCK = PriorityLock()
PRIO_TEACHER, PRIO_BACKGROUND = 0, 1


class UrlVideoChat(LlmChat):
    """Sends a public video URL (e.g. YouTube) for Gemini to watch directly."""
    video_url: Optional[str] = None

    async def _add_user_message(self, messages, message):
        messages.append({"role": "user", "content": [
            {"type": "file", "file": {"file_id": self.video_url, "format": "video/mp4"}},
            {"type": "text", "text": message.text}]})
        await self._save_messages(messages)


async def ask_llm(system_message: str, session_id: str, prompt: str, prio: int = PRIO_BACKGROUND, attempts: int = 5,
                  video_url: Optional[str] = None, video_path: Optional[str] = None) -> str:
    last = None
    for i in range(attempts):
        try:
            await LLM_LOCK.acquire(prio)
            try:
                cls = UrlVideoChat if video_url else LlmChat
                chat = cls(api_key=EMERGENT_LLM_KEY, session_id=f"{session_id}-{uuid.uuid4().hex[:6]}",
                           system_message=system_message).with_model("gemini", "gemini-3.1-pro-preview")
                if video_url:
                    chat.video_url = video_url
                files = [FileContentWithMimeType(mime_type="video/mp4", file_path=video_path)] if video_path else None
                return await chat.send_message(UserMessage(text=prompt, file_contents=files))
            finally:
                LLM_LOCK.release()
        except Exception as e:  # noqa: BLE001
            last = e
            logger.warning(f"LLM call failed (attempt {i + 1}/{attempts}): {str(e)[:150]}")
            if i < attempts - 1:
                await asyncio.sleep(min(30, 3 * 2 ** i))
    raise last


# ---------- AI grading ----------
SYSTEM_PROMPT = """Kamu adalah Asisten Guru Seni Musik SMA yang ahli dalam menganalisis konten video (visual, audio, dan teks) serta objektif dalam memberikan nilai. Tugasmu adalah mengevaluasi video tugas siswa yang berjudul 'Creative Video Project: Musik di Sekitar Kita'.
ATURAN UTAMA (ERROR HANDLING): Jika kamu menerima pesan bahwa video tidak dapat diekstrak, atau jika kamu tidak bisa mengakses/menonton isi video tersebut karena pembatasan privasi tautan (private link), JANGAN mengarang nilai. Langsung berikan format berikut:
ai_score: 0
ai_letter_grade: "N/A"
ai_strengths: "-"
ai_weaknesses: "Sistem tidak dapat menonton video karena tautan diprivasi atau diblokir platform. Silakan klik tautan dan nilai secara manual."
RUBRIK PENILAIAN (Jika video berhasil diakses/ditonton): Nilai video secara keseluruhan dalam skala 0-100 berdasarkan 3 kriteria berikut:
Content & Context (Bobot 50%): Analisis ISI VIDEO (visual dan audio) beserta caption-nya. Apakah video tersebut menjelaskan fungsi musik di dunia nyata secara akurat, kreatif, dan menarik? Apakah ada contoh nyata berupa audio atau visual yang mendukung penjelasan tersebut?
Delivery & Subtitles (Bobot 30%): Bagaimana cara penyampaian siswa di dalam video? Apakah komunikatif, jelas, dan percaya diri? Periksa juga apakah terdapat subtitle atau teks di layar video untuk membantu penyampaian.
Technical & Tagging (Bobot 20%): Apakah video berorientasi vertikal (9:16)? Apakah terdapat hashtag #FungsiMusik dan mention @Mr. Ocha di dalam caption atau teks video?
FORMAT OUTPUT WAJIB (JSON MURNI): Kamu hanya boleh menjawab dengan format JSON yang terstruktur, menggunakan bahasa Indonesia yang baku dan membangun. Jangan tambahkan teks apa pun di luar JSON ini: { "ai_score": [Angka 0-100], "ai_letter_grade": "[A untuk >90, B untuk 80-89, C untuk 70-79, D untuk <70]", "ai_strengths": "[Tuliskan 1-2 kalimat spesifik mengenai kekuatan atau hal positif yang ditemukan dari isi video dan penyampaiannya]", "ai_weaknesses": "[Tuliskan 1-2 kalimat spesifik mengenai kekurangan video dan saran perbaikan untuk siswa]" }"""

HASHTAG_RE = re.compile(r"#[\w\u00C0-\u024F]+", re.UNICODE)
MENTION_RE = re.compile(r"@[\w.]+(?:\s(?:Ocha))?", re.UNICODE)


def metadata_text(meta: dict) -> str:
    """Build the text payload for the AI. Returns EXTRACTION_FAILED if nothing readable was extracted."""
    if not has_readable_text(meta):
        return EXTRACTION_FAILED
    page = meta.get("page", {})
    title = meta.get("oembed", {}).get("title") or page.get("title") or "-"
    caption = page.get("full_description") or page.get("description") or "-"
    blob = " ".join(str(x) for x in (title, caption, page.get("keywords", "")))
    hashtags = sorted(set(HASHTAG_RE.findall(blob)))
    mentions = sorted(set(m.strip() for m in MENTION_RE.findall(blob)))
    lines = [f"Platform: {meta.get('platform')}", f"Judul Video: {title}",
             f"Akun Pengunggah: {meta.get('oembed', {}).get('author_name', '-')}",
             f"Caption/Deskripsi: {caption}", f"Hashtag: {', '.join(hashtags) or '-'}",
             f"Mention: {', '.join(mentions) or '-'}"]
    if page.get("keywords"):
        lines.append(f"Keywords: {page['keywords']}")
    if page.get("duration_seconds"):
        lines.append(f"Durasi: {page['duration_seconds']} detik")
    if meta.get("url_hint"):
        lines.append(f"Petunjuk URL: {meta['url_hint']}")
    return "\n".join(lines)


def parse_json(text: str) -> dict:
    text = text.strip()
    m = re.search(r"\{.*\}", text, re.S)
    return json.loads(m.group(0) if m else text)


def clamp(v, default=0.0) -> float:
    try:
        return round(max(0.0, min(100.0, float(v))), 1)
    except (TypeError, ValueError):
        return default


def download_video(url: str, folder: str) -> Optional[str]:
    opts = {"outtmpl": f"{folder}/video.%(ext)s", "format": "b[ext=mp4][height<=720]/b[height<=720]/b",
            "max_filesize": MAX_VIDEO_BYTES, "quiet": True, "no_warnings": True, "noplaylist": True,
            "socket_timeout": 20, "ffmpeg_location": imageio_ffmpeg.get_ffmpeg_exe(),
            "merge_output_format": "mp4"}
    try:
        with yt_dlp.YoutubeDL(opts) as ydl:
            ydl.download([url])
    except Exception as e:  # noqa: BLE001
        logger.info(f"Video download failed: {str(e)[:200]}")
        return None
    files = [p for p in Path(folder).iterdir() if p.is_file() and p.stat().st_size > 0]
    return str(files[0]) if files else None


def caption_text(meta: dict) -> str:
    return metadata_text(meta) if has_readable_text(meta) else "Caption/metadata tidak tersedia."


async def grade_submission(sub_id: str):
    sub = await db.submissions.find_one({"id": sub_id}, {"_id": 0})
    if not sub:
        return
    await db.submissions.update_one({"id": sub_id}, {"$set": {"status": "processing", "error": None}})
    folder = tempfile.mkdtemp(prefix="vibesmai-")
    try:
        meta = await fetch_metadata(sub["video_link"], sub["platform"])
        video_url = video_path = None
        if sub["platform"] == "youtube":
            video_url = sub["video_link"] if meta.get("oembed") else None
        else:
            video_path = await asyncio.wait_for(asyncio.to_thread(download_video, sub["video_link"], folder), 180)
        extraction_ok = bool(video_url or video_path)
        payload = caption_text(meta) if extraction_ok else EXTRACTION_FAILED
        res = {}
        if extraction_ok:
            prompt = (f"Siswa: {sub['full_name']} (Kelas {sub['class_name']}, Absen {sub['attendance_number']})\n"
                      f"Link video: {sub['video_link']}\n\nVideo tugas terlampir. Caption/metadata video:\n{payload}")
            reply = await ask_llm(SYSTEM_PROMPT, f"grade-{sub_id}", prompt, PRIO_BACKGROUND, attempts=5,
                                  video_url=video_url, video_path=video_path)
            res = parse_json(reply)
        if not extraction_ok or str(res.get("ai_letter_grade", "")).upper() == "N/A":
            extraction_ok = False
            score, grade, strengths, weaknesses = 0.0, "N/A", "-", PRIVACY_WEAKNESS
        else:
            score = clamp(res.get("ai_score"))
            grade = letter_grade(score)
            strengths = res.get("ai_strengths", "")
            weaknesses = res.get("ai_weaknesses", "")
        ai = {"ai_score": score, "ai_letter_grade": grade, "ai_strengths": strengths, "ai_weaknesses": weaknesses,
              "raw_letter_grade": res.get("ai_letter_grade"), "video_mode": "url" if video_url else "file" if video_path else None,
              "graded_at": now_iso()}
        await db.submissions.update_one({"id": sub_id}, {"$set": {
            "status": "draft", "metadata": meta, "ai_input": payload, "extraction_ok": extraction_ok, "ai": ai,
            "ai_score": score, "ai_letter_grade": grade, "ai_strengths": strengths, "ai_weaknesses": weaknesses,
            "content_score": None, "delivery_score": None, "technical_score": None,
            "final_score": score, "final_grade": grade, "manually_edited": False, "updated_at": now_iso()}})
    except Exception as e:
        logger.exception("AI grading failed")
        await db.submissions.update_one({"id": sub_id}, {"$set": {"status": "failed", "error": str(e)[:300]}})
    finally:
        shutil.rmtree(folder, ignore_errors=True)


# ---------- Student feedback ----------
FEEDBACK_PROMPT = """Kamu adalah Guru Seni Musik SMA (Mr. Ocha) yang hangat dan memotivasi. Tugasmu menulis komentar singkat untuk siswa tentang tugas video 'Musik di Sekitar Kita'.
ATURAN:
- Tulis dalam Bahasa Indonesia, 2-3 kalimat saja, nada hangat, positif, dan memotivasi.
- Sapa siswa dengan nama panggilannya (nama depan) di awal kalimat.
- Sebutkan satu kelebihan utama, lalu satu saran perbaikan yang konkret dan membangun.
- HANYA gunakan informasi dari data Kelebihan/Kekurangan/Catatan guru. JANGAN mengarang detail isi video yang tidak disebutkan. Jika data tersebut kosong atau '-', tulis komentar umum yang mengapresiasi usaha siswa dan mengingatkan ketentuan tugas (penjelasan fungsi musik + contoh nyata, subtitle, #FungsiMusik dan mention @Mr. Ocha).
- Jangan menyebut AI, sistem, metadata, privasi, atau angka skor. Jangan memakai emoji berlebihan (maksimal 1).
- Balas HANYA dengan teks komentar, tanpa judul, tanpa tanda kutip, tanpa markdown."""


async def generate_feedback(sub: dict, strengths: Optional[str] = None, weaknesses: Optional[str] = None,
                            score: Optional[float] = None, prio: int = PRIO_BACKGROUND, attempts: int = 4) -> str:
    strengths = strengths if strengths is not None else (sub.get("ai_strengths") or "-")
    weaknesses = weaknesses if weaknesses is not None else (sub.get("ai_weaknesses") or "-")
    score = score if score is not None else sub.get("final_score")
    if PRIVACY_WEAKNESS.lower() in weaknesses.lower():
        weaknesses = "-"
    if strengths.strip().lower().startswith("tidak ada"):
        strengths = "-"
    prompt = (f"Nama siswa: {sub['full_name']}\nNilai akhir: {score if score is not None else '-'} "
              f"(grade {letter_grade(score) if score is not None else '-'})\n"
              f"Kelebihan: {strengths}\nKekurangan & saran: {weaknesses}\n"
              f"Catatan guru: {sub.get('teacher_notes') or '-'}")
    reply = await ask_llm(FEEDBACK_PROMPT, f"feedback-{sub['id']}", prompt, prio, attempts=attempts)
    return reply.strip().strip('"').strip()[:2000]


async def feedback_job(sub_id: str, overrides: Optional[dict] = None, only_if_empty: bool = False,
                       prio: int = PRIO_TEACHER):
    """Background: (re)generate the student comment and store it. Status tracked in feedback_status."""
    s = await db.submissions.find_one({"id": sub_id}, {"_id": 0})
    if not s:
        return
    if only_if_empty and (s.get("status") != "final" or (s.get("student_feedback") or "").strip()):
        if s.get("feedback_status") == "generating":
            await db.submissions.update_one({"id": sub_id}, {"$set": {"feedback_status": "ready"}})
        return
    o = overrides or {}
    try:
        fb = await generate_feedback(s, o.get("ai_strengths"), o.get("ai_weaknesses"), o.get("final_score"), prio=prio)
        query = {"id": sub_id}
        if only_if_empty:
            query["$or"] = [{"student_feedback": {"$in": [None, ""]}}, {"student_feedback": {"$exists": False}}]
        await db.submissions.update_one(query, {"$set": {"student_feedback": fb, "feedback_status": "ready",
                                                         "feedback_generated_at": now_iso()}})
        await db.submissions.update_one({"id": sub_id, "feedback_status": "generating"}, {"$set": {"feedback_status": "ready"}})
    except Exception as e:
        logger.exception("Feedback generation failed")
        await db.submissions.update_one({"id": sub_id}, {"$set": {"feedback_status": "failed",
                                                                  "feedback_error": str(e)[:200]}})


# ---------- Public ----------
async def results_public() -> bool:
    doc = await db.settings.find_one({"key": "results_public"}, {"_id": 0})
    return bool(doc and doc.get("value"))


@api.get("/")
async def root():
    return {"message": "VIBESMAI API"}


@api.post("/submissions")
async def create_submission(data: SubmissionIn, bg: BackgroundTasks):
    url = (data.video_link or data.video_url or "").strip()
    if not url:
        raise HTTPException(422, "Video link is required")
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    platform = detect_platform(url)
    if not platform:
        raise HTTPException(422, "Link must be from TikTok, Instagram, Facebook, or YouTube")
    doc = {"id": str(uuid.uuid4()), "full_name": data.full_name, "class_name": data.class_name,
           "attendance_number": data.attendance_number, "video_link": url, "platform": platform,
           "status": "pending", "ai_score": None, "ai_letter_grade": None, "ai_strengths": None,
           "ai_weaknesses": None, "final_score": None, "final_grade": None,
           "manually_edited": False, "teacher_notes": "", "created_at": now_iso()}
    await db.submissions.insert_one(doc)
    bg.add_task(grade_submission, doc["id"])
    return {"ok": True, "message": "Your video assignment link has been successfully submitted / Tugas link video Anda berhasil dikirimkan."}


@api.get("/public/settings")
async def public_settings():
    return {"results_public": await results_public()}


@api.get("/public/results")
async def public_results(class_name: str, attendance_number: int):
    if not await results_public():
        raise HTTPException(403, "Results page is disabled")
    subs = await db.submissions.find({"class_name": class_name, "attendance_number": attendance_number},
                                     {"_id": 0}).sort("created_at", -1).to_list(50)
    out = []
    for s in subs:
        final = s.get("status") == "final"
        out.append({"full_name": s["full_name"], "class_name": s["class_name"],
                    "attendance_number": s["attendance_number"], "platform": s["platform"],
                    "created_at": s["created_at"], "is_final": final,
                    "final_score": s.get("final_score") if final else None,
                    "final_grade": s.get("final_grade") if final else None,
                    "student_feedback": (s.get("student_feedback") or None) if final else None})
    return out


# ---------- Admin ----------
def build_query(class_name, platform, status, q):
    query = {}
    if class_name:
        query["class_name"] = class_name
    if platform:
        query["platform"] = platform
    if status == "pending":
        query["status"] = {"$in": ["pending", "processing"]}
    elif status:
        query["status"] = status
    if q:
        query["full_name"] = {"$regex": re.escape(q), "$options": "i"}
    return query


def class_sort_key(s):
    return (int(s["class_name"].split()[-1]), s["attendance_number"], s["created_at"])


@api.get("/admin/submissions")
async def list_submissions(class_name: Optional[str] = None, platform: Optional[str] = None,
                           status: Optional[str] = None, q: Optional[str] = None,
                           _: dict = Depends(require_admin)):
    return await db.submissions.find(build_query(class_name, platform, status, q),
                                     {"_id": 0, "metadata": 0, "ai_input": 0}).sort("created_at", -1).to_list(5000)


@api.get("/admin/stats")
async def stats(_: dict = Depends(require_admin)):
    subs = await db.submissions.find({}, {"_id": 0, "status": 1, "final_score": 1, "class_name": 1}).to_list(10000)
    scored = [s for s in subs if s.get("final_score") is not None]
    per_class = {c: 0 for c in CLASSES}
    for s in subs:
        per_class[s["class_name"]] = per_class.get(s["class_name"], 0) + 1
    count = lambda *st: sum(1 for s in subs if s["status"] in st)  # noqa: E731
    return {"total": len(subs), "draft": count("draft"), "final": count("final"),
            "pending": count("pending", "processing"), "failed": count("failed"),
            "avg_score": round(sum(s["final_score"] for s in scored) / len(scored), 1) if scored else None,
            "per_class": per_class}


@api.get("/admin/submissions/{sub_id}")
async def get_submission(sub_id: str, _: dict = Depends(require_admin)):
    s = await db.submissions.find_one({"id": sub_id}, {"_id": 0})
    if not s:
        raise HTTPException(404, "Not found")
    return s


@api.patch("/admin/submissions/{sub_id}")
async def update_submission(sub_id: str, data: SubmissionUpdate, bg: BackgroundTasks, _: dict = Depends(require_admin)):
    s = await db.submissions.find_one({"id": sub_id}, {"_id": 0})
    if not s:
        raise HTTPException(404, "Not found")
    upd = data.model_dump(exclude_none=True)
    score_keys = {"content_score", "delivery_score", "technical_score"}
    if "final_score" in upd:
        upd["final_score"] = round(upd["final_score"], 1)
    elif score_keys & upd.keys():
        merged = {k: upd.get(k, s.get(k)) for k in score_keys}
        if any(v is None for v in merged.values()):
            raise HTTPException(422, "All three rubric scores are required")
        upd["final_score"] = weighted(merged["content_score"], merged["delivery_score"], merged["technical_score"])
    if "final_score" in upd:
        upd["final_grade"] = letter_grade(upd["final_score"])
    new_status = upd.get("status")
    if new_status == "final":
        if (upd.get("final_score") if "final_score" in upd else s.get("final_score")) is None:
            raise HTTPException(422, "Cannot finalize without a score")
        upd["finalized_at"] = now_iso()
        feedback = upd["student_feedback"] if "student_feedback" in upd else s.get("student_feedback")
        if not (feedback or "").strip():
            upd["feedback_status"] = "generating"
            bg.add_task(feedback_job, sub_id, None, True, PRIO_TEACHER)
    elif new_status == "draft" and s.get("status") in ("pending", "processing", "failed") and "final_score" not in upd:
        upd.pop("status")
    if set(upd) - {"status", "teacher_notes", "finalized_at", "student_feedback", "feedback_status"}:
        upd["manually_edited"] = True
    upd["updated_at"] = now_iso()
    await db.submissions.update_one({"id": sub_id}, {"$set": upd})
    return await db.submissions.find_one({"id": sub_id}, {"_id": 0})


@api.post("/admin/submissions/{sub_id}/regrade")
async def regrade(sub_id: str, bg: BackgroundTasks, _: dict = Depends(require_admin)):
    res = await db.submissions.update_one({"id": sub_id}, {"$set": {"status": "pending"}})
    if not res.matched_count:
        raise HTTPException(404, "Not found")
    bg.add_task(grade_submission, sub_id)
    return {"ok": True}


@api.post("/admin/submissions/{sub_id}/feedback")
async def regenerate_feedback(sub_id: str, data: FeedbackIn, bg: BackgroundTasks, _: dict = Depends(require_admin)):
    res = await db.submissions.update_one({"id": sub_id}, {"$set": {"feedback_status": "generating", "feedback_error": None}})
    if not res.matched_count:
        raise HTTPException(404, "Not found")
    bg.add_task(feedback_job, sub_id, data.model_dump(exclude_none=True), False, PRIO_TEACHER)
    return {"feedback_status": "generating"}


@api.delete("/admin/submissions/{sub_id}")
async def delete_submission(sub_id: str, _: dict = Depends(require_admin)):
    res = await db.submissions.delete_one({"id": sub_id})
    if not res.deleted_count:
        raise HTTPException(404, "Not found")
    return {"ok": True}


@api.post("/admin/bulk-status")
async def bulk_status(data: BulkStatus, bg: BackgroundTasks, _: dict = Depends(require_admin)):
    query = {"id": {"$in": data.ids}, "final_score": {"$ne": None}}
    upd = {"status": data.status, "updated_at": now_iso()}
    if data.status == "final":
        upd["finalized_at"] = now_iso()
    res = await db.submissions.update_many(query, {"$set": upd})
    if data.status == "final":
        targets = await db.submissions.find({"id": {"$in": data.ids}, "status": "final",
                                             "$or": [{"student_feedback": {"$in": [None, ""]}},
                                                     {"student_feedback": {"$exists": False}}]},
                                            {"_id": 0, "id": 1}).to_list(5000)
        if targets:
            await db.submissions.update_many({"id": {"$in": [t["id"] for t in targets]}},
                                             {"$set": {"feedback_status": "generating"}})
        for t in targets:
            bg.add_task(feedback_job, t["id"], None, True, PRIO_BACKGROUND)
    return {"updated": res.modified_count}


@api.get("/admin/settings")
async def get_settings(user: dict = Depends(require_admin)):
    return {"results_public": await results_public(), "admin_email": await get_admin_email()}


@api.put("/admin/settings")
async def put_settings(data: SettingsIn, _: dict = Depends(require_admin)):
    await db.settings.update_one({"key": "results_public"}, {"$set": {"value": data.results_public}}, upsert=True)
    return {"results_public": data.results_public}


STATUS_LABEL = {"pending": "Diproses", "processing": "Diproses", "draft": "Draft", "final": "Final", "failed": "Gagal"}
EXPORT_COLS = [("Kelas", "class_name"), ("No. Absen", "attendance_number"), ("Nama Siswa", "full_name"),
               ("Platform", "platform"), ("Link Video", "video_link"), ("Waktu Kirim", "created_at"),
               ("Status", "status"), ("Skor AI", "ai_score"), ("Grade AI", "ai_letter_grade"),
               ("Content & Context (50%)", "content_score"), ("Delivery & Subtitles (30%)", "delivery_score"),
               ("Technical & Tagging (20%)", "technical_score"), ("Nilai Akhir", "final_score"),
               ("Letter Grade", "final_grade"), ("Diedit Manual", "manually_edited"),
               ("Kelebihan (AI)", "ai_strengths"), ("Kekurangan & Saran (AI)", "ai_weaknesses"),
               ("Komentar untuk Siswa", "student_feedback"), ("Catatan Guru", "teacher_notes")]
WIDE_COLS = {"video_link", "ai_strengths", "ai_weaknesses", "teacher_notes", "student_feedback"}


def export_cell(s: dict, k: str):
    v = s.get(k)
    if k == "manually_edited":
        return "Ya" if v else "Tidak"
    if k == "status":
        return STATUS_LABEL.get(v, v or "")
    return "" if v is None else v


@api.get("/admin/export")
async def export(format: str = Query("csv", pattern="^(csv|xlsx)$"), class_name: Optional[str] = None,
                 _: dict = Depends(require_admin)):
    subs = await db.submissions.find(build_query(class_name, None, None, None),
                                     {"_id": 0, "metadata": 0, "ai_input": 0}).to_list(10000)
    subs.sort(key=class_sort_key)
    rows = [[export_cell(s, k) for _, k in EXPORT_COLS] for s in subs]
    headers = [h for h, _ in EXPORT_COLS]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    suffix = f"_kelas_{class_name.replace(' ', '_')}" if class_name else ""
    fname = f"vibesmai_rekap_nilai{suffix}_{stamp}"
    if format == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")
        w = csv.writer(buf)
        w.writerow(headers)
        w.writerows(rows)
        return Response(buf.getvalue(), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})
    wb = Workbook()
    wb.remove(wb.active)
    if class_name:
        groups = [(f"Kelas {class_name}", list(range(len(subs))))]
    else:
        groups = [("Semua Kelas", list(range(len(subs))))] + \
                 [(c, [i for i, s in enumerate(subs) if s["class_name"] == c]) for c in CLASSES]
    for idx, (title, items) in enumerate(groups):
        if idx > 0 and not items:
            continue
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="111827")
        for i in items:
            ws.append(rows[i])
        for i, (h, k) in enumerate(EXPORT_COLS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 40 if k in WIDE_COLS else max(12, len(h) + 2)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})


async def migrate_legacy():
    """Convert documents from the previous schema (video_url/submitted_at/grade/published) to the new one."""
    async for s in db.submissions.find({"video_url": {"$exists": True}}, {"_id": 0}):
        status = s.get("status")
        if status == "graded":
            status = "final" if s.get("published") else "draft"
        grade = s.get("grade")
        grade = "D" if grade == "N/A" else grade
        ai = s.get("ai") or {}
        upd = {"video_link": s["video_url"], "created_at": s.get("submitted_at") or now_iso(), "status": status,
               "ai_score": ai.get("final_score", s.get("final_score")),
               "ai_letter_grade": "D" if ai.get("grade") == "N/A" else ai.get("grade", grade),
               "ai_strengths": s.get("strengths"), "ai_weaknesses": s.get("weaknesses"), "final_grade": grade}
        await db.submissions.update_one({"id": s["id"]}, {"$set": upd, "$unset": {
            "video_url": "", "submitted_at": "", "grade": "", "strengths": "", "weaknesses": "", "published": ""}})


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True,
                   allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
                   allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
async def startup():
    await db.submissions.create_index("id", unique=True)
    await db.submissions.create_index([("class_name", 1), ("attendance_number", 1)])
    await migrate_legacy()
    await db.users.create_index("email", unique=True)
    await db.user_sessions.create_index("session_token")
    if ADMIN_EMAIL:
        await db.settings.update_one({"key": "admin"}, {"$setOnInsert": {"key": "admin", "email": ADMIN_EMAIL}}, upsert=True)
    # Recover submissions interrupted by a restart
    stuck = await db.submissions.find({"status": {"$in": ["pending", "processing"]}}, {"_id": 0, "id": 1}).to_list(500)
    for s in stuck:
        asyncio.create_task(grade_submission(s["id"]))
    fb_stuck = await db.submissions.find({"feedback_status": "generating"}, {"_id": 0, "id": 1}).to_list(500)
    for s in fb_stuck:
        asyncio.create_task(feedback_job(s["id"], None, True, PRIO_BACKGROUND))


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
